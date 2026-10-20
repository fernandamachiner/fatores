# Confere, célula a célula, o cadastro metodologia/metodologia.js contra a
# planilha oficial (metodologia/planilha/Fatores_organizacao.xlsx).
# A leitura da planilha é feita aqui de forma independente do gerador.
#
# Uso, a partir da pasta do painel:
#   python3 testes/conferir_metodologia_planilha.py
import json, sys
import openpyxl

M = json.loads(open("metodologia/metodologia.js", encoding="utf-8").read().split("window.METODOLOGIA = ", 1)[1].rstrip().rstrip(";"))
cad = [(e, f) for e in M["eixos"] for f in e["fatores"]]
ws = openpyxl.load_workbook("metodologia/planilha/Fatores_organizacao.xlsx", data_only=True)["Consolidado"]
linhas = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r[:8])]
erros = []
if len(linhas) != len(cad):
    erros.append("planilha com %d fatores, cadastro com %d" % (len(linhas), len(cad)))
celulas = 0
for n, (r, (e, f)) in enumerate(zip(linhas, cad), start=2):
    esperado = [e["id_original"], e["titulo"], f["fator"], f["icone"], f["definicao"], f["ods"], f["nome_ods"], f["id"]]
    for c, (v, x) in enumerate(zip(r[:8], esperado)):
        celulas += 1
        v = v.strip() if isinstance(v, str) else v
        if c == 5:
            v = int(v)
        if v != x:
            erros.append("linha %d, coluna %s: planilha %r, cadastro %r" % (n, ws.cell(1, c + 1).value, v, x))
print("Planilha: %d fatores | cadastro: %d fatores | células conferidas: %d | divergências: %d" % (len(linhas), len(cad), celulas, len(erros)))
for x in erros[:30]:
    print("ERRO  " + x)
sys.exit(1 if erros else 0)
