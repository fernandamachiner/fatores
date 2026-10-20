# Monta o documento do TERRITÓRIO DEMONSTRATIVO (SIMULADO) a partir do modelo
# territorios/modelos/modelo-diagnostico.docx, como faria uma equipe ao
# preencher o modelo. Todos os textos são fictícios e identificados como
# simulados; o território não existe e não deve ser publicado.
#
# Preenche 12 fatores (2 por eixo) com textos simulados, declara um fator sem
# diagnóstico com motivo e os demais como "Sem diagnóstico disponível.", e
# apaga as linhas de instrução do modelo.
#
# Uso, a partir da pasta do painel:
#   python3 testes/exemplo-simulado/montar_exemplo.py
import json, re, sys
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

MODELO = "territorios/modelos/modelo-diagnostico.docx"
SAIDA = "testes/exemplo-simulado/Diagnostico_Territorio_Demonstrativo_SIMULADO.docx"
M = json.loads(open("metodologia/metodologia.js", encoding="utf-8").read().split("window.METODOLOGIA = ", 1)[1].rstrip().rstrip(";"))
PREENCHER = {}
for e in M["eixos"]:
    a, b = e["fatores"][0], e["fatores"][1]
    PREENCHER[a["id"]] = ("[DADOS SIMULADOS] Texto fictício de demonstração para o fator %s. Em 2025, a oficina simulada registrou 120 participantes "
                          "(Fonte simulada: Relatório fictício da oficina de 2025). O tema se relaciona ao fator %s.") % (a["fator"], b["id"])
    PREENCHER[b["id"]] = ("[DADOS SIMULADOS] Texto fictício de demonstração para o fator %s, sem valor de análise territorial. "
                          "A Lei Municipal fictícia nº 0000/2024 é citada apenas para demonstrar o registro de normas.") % b["fator"]
MOTIVO = {M["eixos"][0]["fatores"][2]["id"]: "tema não discutido na oficina simulada"}

d = docx.Document(MODELO)
corpo = d.element.body
fator = None
remover = []
for el in list(corpo.iterchildren()):
    tag = el.tag.split("}")[1]
    if tag == "tbl":
        t = Table(el, d)
        ident = t.rows[0].cells[0].text.strip()
        if len(t.rows) == 1 and re.match(r"^E[1-6]-F\d{2}$", ident):
            fator = ident
        else:  # tabela de abertura do eixo
            t.cell(0, 0).text = "Contradição central: [DADOS SIMULADOS] contradição fictícia de demonstração."
            t.cell(1, 0).text = "Foco pedagógico: [DADOS SIMULADOS] foco fictício de demonstração."
        continue
    if tag != "p":
        continue
    p = Paragraph(el, d)
    txt = p.text
    if "[PREENCHER" not in txt:
        continue
    if txt.startswith("Diagnóstico Situacional"):
        for r in p.runs: r.text = ""
        p.runs[0].text = "Diagnóstico Situacional - Território demonstrativo (SIMULADO)"
    elif "diagnóstico situacional do fator" in txt:
        for r in p.runs: r.text = ""
        if fator in PREENCHER:
            p.runs[0].text = PREENCHER[fator]
        elif fator in MOTIVO:
            p.runs[0].text = "Sem diagnóstico disponível: %s." % MOTIVO[fator]
        else:
            p.runs[0].text = "Sem diagnóstico disponível."
        p.runs[0].italic = False
    elif "responsáveis pela elaboração" in txt:
        for r in p.runs: r.text = ""
        p.runs[0].text = "Documento SIMULADO, produzido apenas para demonstrar o procedimento de cadastro. Período de referência fictício: 2024 a 2025."
    else:
        remover.append(el)
for el in remover:
    corpo.remove(el)
d.save(SAIDA)
restam = sum(1 for p in docx.Document(SAIDA).paragraphs if "[PREENCHER" in p.text)
print("%s: %d fatores com texto simulado, %d declarado com motivo, linhas de instrução restantes: %d" % (SAIDA, len(PREENCHER), len(MOTIVO), restam))
sys.exit(1 if restam else 0)
