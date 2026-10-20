# Gera o modelo de documento de diagnóstico situacional no formato padrão
# da plataforma (territorios/modelos/modelo-diagnostico.docx), a partir do
# cadastro metodológico vigente. O modelo já traz, para os 150 fatores, a
# tabela de identificação [ID | nome | ODS] e a definição de referência,
# idênticas ao cadastro; a equipe do território escreve apenas os textos.
#
# Os trechos a preencher começam por [PREENCHER. O gerador do território
# recusa o documento enquanto algum desses trechos permanecer.
#
# Gere de novo sempre que a metodologia mudar (nome, ODS ou definição de
# algum fator), para que o modelo continue idêntico ao cadastro.
#
# Uso, a partir da pasta do painel:
#   python3 territorios/ferramentas/gerar_modelo_documento.py
import json
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SAIDA = "territorios/modelos/modelo-diagnostico.docx"


def carregar_metodologia():
    txt = open("metodologia/metodologia.js", encoding="utf-8").read()
    return json.loads(txt.split("window.METODOLOGIA = ", 1)[1].rstrip().rstrip(";"))


def instrucao(d, texto):
    p = d.add_paragraph()
    r = p.add_run(texto)
    r.italic = True
    r.font.color.rgb = RGBColor(0xA0, 0x40, 0x00)
    return p


def main():
    M = carregar_metodologia()
    d = docx.Document()
    estilo = d.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # capa
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Diagnóstico Situacional - [PREENCHER: nome da localidade e UF]")
    r.bold = True
    r.font.size = Pt(16)
    d.add_paragraph("Organizado pela %s v%s (%d fatores)" % (M["nome"], M["versao"], sum(len(e["fatores"]) for e in M["eixos"])))
    instrucao(d, "[PREENCHER: responsáveis pela elaboração, data e período de referência dos dados, por exemplo 2020 a 2026]")
    instrucao(d, "[PREENCHER: apague esta e todas as linhas que começam por PREENCHER depois de escrever o conteúdo. "
                 "Não altere as tabelas de identificação dos fatores nem as linhas que começam por Definição de referência. "
                 "Para um fator sem informação, escreva apenas: Sem diagnóstico disponível. ou Sem diagnóstico disponível: motivo.]")

    d.add_heading("Resumo executivo", level=1)
    instrucao(d, "[PREENCHER: síntese do diagnóstico da localidade, opcional. Se não houver, apague o título e esta linha]")

    for e in M["eixos"]:
        d.add_heading("Eixo %d. %s" % (e["ordem"], e["titulo"]), level=1)
        tab = d.add_table(rows=2, cols=1)
        tab.style = "Table Grid"
        tab.cell(0, 0).text = "Contradição central: [PREENCHER: opcional; se não houver, apague a tabela inteira]"
        tab.cell(1, 0).text = "Foco pedagógico: [PREENCHER: opcional]"
        instrucao(d, "[PREENCHER: introdução do eixo, opcional]")
        for f in e["fatores"]:
            tf = d.add_table(rows=1, cols=3)
            tf.style = "Table Grid"
            tf.cell(0, 0).text = f["id"]
            tf.cell(0, 1).text = f["fator"]
            tf.cell(0, 2).text = "ODS %d" % f["ods"]
            for c in tf.rows[0].cells:
                for run in c.paragraphs[0].runs:
                    run.bold = True
            d.add_paragraph("Definição de referência: " + f["definicao"])
            instrucao(d, "[PREENCHER: diagnóstico situacional do fator %s na localidade, citando a fonte e o ano de cada dado, "
                         "ou apenas Sem diagnóstico disponível.]" % f["id"])

    d.add_heading("Síntese estratégica", level=1)
    instrucao(d, "[PREENCHER: opcional. Pode usar títulos de nível 2, como Nós críticos, Potencialidades e Prioridades]")
    d.add_heading("Notas metodológicas e fontes", level=1)
    instrucao(d, "[PREENCHER: como o diagnóstico foi produzido, período de referência e lista de fontes consultadas]")
    d.save(SAIDA)
    print("%s: %d eixos e %d fatores (metodologia v%s)" % (SAIDA, len(M["eixos"]), sum(len(e["fatores"]) for e in M["eixos"]), M["versao"]))


if __name__ == "__main__":
    main()
